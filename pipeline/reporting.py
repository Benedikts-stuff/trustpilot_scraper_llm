# reporting.py
import pandas as pd
import os
from openpyxl import load_workbook
import matplotlib.pyplot as plt
from wordcloud import WordCloud
import seaborn as sns
from nltk.corpus import stopwords


def write_to_excel(filepath, sheet_name, df):
    """Schreibt einen DataFrame in eine Excel-Datei, überschreibt Sheet, wenn vorhanden."""
    if os.path.exists(filepath):
        try:
            book = load_workbook(filepath)
            if sheet_name in book.sheetnames:
                # Sheet löschen, wenn es existiert
                sheet = book[sheet_name]
                book.remove(sheet)

            # Pandas ExcelWriter im 'append'-Modus verwenden, aber da wir
            # das Sheet gelöscht haben, wird es neu erstellt.
            with pd.ExcelWriter(filepath, engine='openpyxl', mode='a') as writer:
                writer.book = book
                df.to_excel(writer, sheet_name=sheet_name, index=False)
        except Exception as e:
            print(f"Fehler beim Anhängen an Excel: {e}. Erstelle Datei neu.")
            # Fallback: Datei komplett neu schreiben
            with pd.ExcelWriter(filepath, engine='openpyxl', mode='w') as writer:
                df.to_excel(writer, sheet_name=sheet_name, index=False)
    else:
        # Datei existiert noch nicht → neu anlegen
        with pd.ExcelWriter(filepath, engine='openpyxl', mode='w') as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)


def create_wordcloud(text_series, title):
    """Erstellt und zeigt eine Wordcloud aus einer Pandas-Textserie."""
    stopwords_de = set(stopwords.words('german'))

    # Alle Texte zusammenfassen
    full_text = " ".join(text_series)

    # Wörter filtern
    filtered_words = " ".join([
        word for word in full_text.lower().split()
        if word not in stopwords_de and word.isalpha() and len(word) > 2
    ])

    if not filtered_words:
        print("Keine Wörter für Wordcloud gefunden.")
        return None

    wordcloud = WordCloud(width=800, height=400, background_color='white').generate(filtered_words)

    fig, ax = plt.subplots(figsize=(10, 5))
    ax.imshow(wordcloud, interpolation='bilinear')
    ax.axis("off")
    ax.set_title(title)

    # Gibt die Matplotlib-Figur zurück, damit Streamlit sie anzeigen kann
    return fig


def create_sentiment_plot(df_kategorien_sentiment):
    """Erstellt einen gestapelten Balken-Plot für Sentiment-Verteilung."""
    # Diese Funktion nimmt an, dass der DF bereits gruppiert ist,
    # oder wir machen es hier:
    if "Sentiment" not in df_kategorien_sentiment.columns:
        print("DataFrame hat nicht das richtige Format für Sentiment-Plot.")
        return None

    sentiment_counts = df_kategorien_sentiment.groupby(["Kategorie", "Sentiment"]).size().unstack().fillna(0)

    fig, ax = plt.subplots(figsize=(10, 6))
    sentiment_counts.plot(kind="bar", stacked=True, ax=ax)

    ax.set_title("Sentiment-Verteilung pro Kategorie")
    ax.set_ylabel("Anzahl Texte")
    ax.set_xlabel("Kategorie")
    plt.xticks(rotation=45)
    plt.tight_layout()

    return fig